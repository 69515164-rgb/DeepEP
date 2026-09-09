#!/usr/bin/env python3
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parents[1] / "docs" / "NVL72_多QP动态VT调度方案.pptx"

NAVY = RGBColor(10, 24, 48)
BLUE = RGBColor(38, 108, 255)
CYAN = RGBColor(37, 200, 221)
GREEN = RGBColor(33, 181, 115)
ORANGE = RGBColor(255, 153, 51)
RED = RGBColor(225, 72, 72)
WHITE = RGBColor(248, 250, 253)
LIGHT = RGBColor(229, 236, 245)
MID = RGBColor(130, 148, 174)
DARK = RGBColor(28, 39, 57)


def add_text(slide, text, x, y, w, h, size=20, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = "Microsoft YaHei"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line or fill
    return s


def line(slide, x1, y1, x2, y2, color=MID, width=2):
    s = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    s.line.color.rgb = color
    s.line.width = Pt(width)
    return s


def base_slide(prs, title, kicker=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY
    if kicker:
        add_text(slide, kicker.upper(), 0.65, 0.28, 5, 0.25, 9, CYAN, True)
    add_text(slide, title, 0.65, 0.55, 12, 0.55, 27, WHITE, True)
    rect(slide, 0.65, 1.16, 1.0, 0.05, BLUE)
    add_text(slide, str(len(prs.slides)), 12.3, 7.07, 0.4, 0.22, 8, MID,
             align=PP_ALIGN.RIGHT)
    return slide


def pill(slide, text, x, y, w, color=BLUE, size=12):
    rect(slide, x, y, w, 0.38, color, True)
    add_text(slide, text, x, y, w, 0.38, size, WHITE, True, PP_ALIGN.CENTER)


def card(slide, title, body, x, y, w, h, accent=BLUE, title_size=16, body_size=12):
    rect(slide, x, y, w, h, DARK, True, RGBColor(51, 69, 96))
    rect(slide, x, y, 0.07, h, accent)
    add_text(slide, title, x + 0.22, y + 0.12, w - 0.35, 0.35, title_size, WHITE, True)
    add_text(slide, body, x + 0.22, y + 0.53, w - 0.38, h - 0.65,
             body_size, LIGHT, False, valign=MSO_ANCHOR.TOP)


def bullets(slide, items, x, y, w, h, size=16, color=LIGHT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(11)
        p.text = "•  " + p.text
    return box


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 13.333, 0.13, BLUE)
    pill(slide, "架构设计汇报", 0.75, 0.62, 1.65, CYAN)
    add_text(slide, "NVL72 多 QP\n动态 VT 调度方案", 0.75, 1.35, 7.5, 1.55, 34, WHITE, True)
    add_text(slide, "非 Spectrum 以太网络下的流级多路径、PReCCL 式字节重分配与跨 Rail PXN 绕路",
             0.78, 3.22, 7.2, 0.75, 18, LIGHT)
    rect(slide, 8.55, 1.05, 3.9, 4.85, DARK, True, RGBColor(51, 69, 96))
    for i, (name, col) in enumerate([("QP / VT", BLUE), ("RAIL", CYAN), ("PXN", GREEN)]):
        pill(slide, name, 9.15, 1.68 + i * 1.22, 1.25, col, 14)
        line(slide, 10.45, 1.87 + i * 1.22, 11.75, 1.87 + i * 1.22, col, 4)
    add_text(slide, "先换 PATH\n再换 RAIL", 9.12, 5.12, 2.65, 0.52, 18, WHITE, True,
             PP_ALIGN.CENTER)
    add_text(slide, "技术方案 v1.0  |  2026-09", 0.78, 6.72, 4.0, 0.32, 11, MID)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_slide(prs)

    s = base_slide(prs, "结论先行：可行，但目标不是复刻 Spectrum-X", "Executive summary")
    card(s, "可以实现", "保序 RoCE 上的流级多路径\n按运行时拥塞动态改变各 VT 字节份额\n部分故障时隔离慢/坏 QP", 0.7, 1.55, 3.75, 2.05, GREEN)
    card(s, "不能替代", "逐包自适应路由\n乱序接收与重排\n选择性重传", 4.78, 1.55, 3.75, 2.05, RED)
    card(s, "设计原则", "VT = Channel = QP\n每 Rail 保留多个可调 VT\n先 Rail 内换路径，Incast 再跨 Rail", 8.86, 1.55, 3.75, 2.05, BLUE)
    add_text(s, "推荐起点", 0.75, 4.28, 1.5, 0.35, 15, CYAN, True)
    rect(s, 0.75, 4.75, 11.85, 1.25, RGBColor(19, 49, 83), True)
    add_text(s, "8 VT = 4 Rail × 2 QP", 1.05, 4.95, 3.2, 0.48, 23, WHITE, True)
    add_text(s, "L1  Rail 内重分配", 4.42, 4.95, 2.75, 0.48, 18, CYAN, True)
    add_text(s, "→", 7.13, 4.95, 0.35, 0.48, 22, MID, True, PP_ALIGN.CENTER)
    add_text(s, "L2  跨 Rail + PXN", 7.58, 4.95, 3.2, 0.48, 18, GREEN, True)

    s = base_slide(prs, "问题背景：非 Spectrum 环境缺少逐包闭环", "Why now")
    for i, (t, b, c) in enumerate([
        ("Spectrum-X", "交换机逐包选择最空出口\nCX8 乱序落位与重排\n端到端遥测、拥塞控制、重传协同", GREEN),
        ("非 Spectrum", "QP 五元组必须稳定\nECMP 将一条 QP 固定到一条路径\n同一 QP 仍受保序与 Go-Back-N 约束", ORANGE),
        ("生产痛点", "大象流少、路径熵不足\n其它作业可随时与本 QP 冲突\n最慢 VT 决定 collective 完成时间", RED),
    ]):
        card(s, t, b, 0.75 + i * 4.08, 1.55, 3.7, 2.35, c)
    add_text(s, "技术机会", 0.78, 4.35, 1.2, 0.32, 15, CYAN, True)
    bullets(s, [
        "用多个保序 QP 制造流级路径多样性，而不是逐包喷洒",
        "把每个 QP 升格为独立 VT，让通信库能观测、建模和调权",
        "利用可编程交换机做路径校验与队列遥测，不改变 QP 内路径",
    ], 0.82, 4.72, 11.7, 1.8, 17)

    s = base_slide(prs, "NVL72 拓扑：网卡池化边界是 Compute Tray", "Physical topology")
    for tray in range(3):
        x = 0.85 + tray * 4.02
        rect(s, x, 1.45, 3.55, 3.65, DARK, True, RGBColor(51, 69, 96))
        add_text(s, f"Compute Tray {tray}", x + 0.2, 1.58, 2.2, 0.32, 14, WHITE, True)
        add_text(s, "单 OS / PXN 域", x + 2.08, 1.58, 1.15, 0.32, 10, CYAN, True, PP_ALIGN.RIGHT)
        for g in range(4):
            col = [BLUE, CYAN, GREEN, ORANGE][g]
            pill(s, f"GPU{g}", x + 0.25, 2.12 + g * 0.62, 0.78, col, 10)
            line(s, x + 1.04, 2.31 + g * 0.62, x + 1.42, 2.31 + g * 0.62, LIGHT, 2)
            pill(s, f"NIC{g}", x + 1.44, 2.12 + g * 0.62, 0.82, col, 10)
            add_text(s, f"Rail {g}", x + 2.45, 2.12 + g * 0.62, 0.7, 0.38, 10, col, True)
    rect(s, 0.85, 5.47, 11.6, 0.65, RGBColor(19, 49, 83), True)
    add_text(s, "NVLink Switch Fabric：72 GPU 同一 NVLink 域，但 PXN 不能跨 18 个 OS",
             1.1, 5.55, 11.0, 0.45, 15, WHITE, True, PP_ALIGN.CENTER)
    add_text(s, "默认池化：本 Tray 4 张 CX8 / 4 条 Rail；整柜 NIC 池化只保留为故障备份研究项",
             0.88, 6.35, 11.6, 0.45, 15, ORANGE, True, PP_ALIGN.CENTER)

    s = base_slide(prs, "为什么必须 VT = Channel = QP", "Granularity")
    card(s, "错误抽象", "一个 Channel 内藏 2 条 QP\n通信库只看到一个 FIFO stall\n无法知道是哪条 QP 与其它大象流冲突\n也无法单独迁移该 QP 的字节", 0.75, 1.55, 5.55, 3.1, RED)
    card(s, "正确抽象", "每个 QP 独占一个 Channel/VT\n每 VT 有独立 stall、完成率和字节区间\n下一 epoch 可单独改变权重\nRail 与实际 ECMP path 作为标签", 7.03, 1.55, 5.55, 3.1, GREEN)
    line(s, 6.35, 3.08, 6.95, 3.08, CYAN, 5)
    add_text(s, "可观测  →  可决策  →  可执行", 2.6, 5.38, 8.2, 0.55, 23, CYAN, True, PP_ALIGN.CENTER)

    s = base_slide(prs, "推荐映射：8 VT = 4 Rail × 2 QP", "Initial configuration")
    colors = [BLUE, CYAN, GREEN, ORANGE]
    for r in range(4):
        y = 1.5 + r * 1.22
        pill(s, f"RAIL {r}", 0.78, y + 0.23, 1.05, colors[r], 12)
        for q in range(2):
            x = 2.25 + q * 4.45
            idx = r * 2 + q
            rect(s, x, y, 3.72, 0.86, RGBColor(25, 45, 72), True, colors[r])
            add_text(s, f"VT{idx}  =  Channel{idx}  =  QP{idx}", x + 0.2, y + 0.08, 2.7, 0.32, 14, WHITE, True)
            add_text(s, f"sport 固定  |  path {r}{'A' if q == 0 else 'B'}", x + 0.2, y + 0.45, 2.9, 0.24, 10, MID)
        add_text(s, "≠", 6.08, y + 0.18, 0.35, 0.45, 20, colors[r], True, PP_ALIGN.CENTER)
    card(s, "初始化约束", "用交换机流表或主动探测确认 path_id；同 Rail 两条 QP 应落在不同上行。无法正交时，不把它们计作两份独立容量。", 0.82, 6.38, 11.65, 0.72, BLUE, 13, 11)

    s = base_slide(prs, "两级控制：先换 PATH，再换 RAIL", "Control hierarchy")
    card(s, "L1：Rail 内调节", "触发：同 Rail 部分 VT 慢、部分正常\n判因：某条 leaf→spine 上行发生 ECMP 冲突\n动作：只在同 Rail 的 QP/VT 间迁移权重\n成本：最低，不增加 NVLink 绕路", 0.8, 1.5, 5.55, 3.45, CYAN)
    card(s, "L2：跨 Rail 绕路", "触发：同 Rail 所有 VT 同时恶化\n判因：目的 leaf 下行或 NIC 入口 incast\n动作：迁移到冷 Rail 的 VT\n路径：其它 NIC 入站 → 本 Tray NVLink/PXN → 目标 GPU", 6.98, 1.5, 5.55, 3.45, GREEN)
    rect(s, 0.85, 5.48, 11.6, 0.75, RGBColor(19, 49, 83), True)
    add_text(s, "避免过度反应：单 QP 冲突不应直接触发跨 Rail；跨 Rail 是处理最后一跳共享瓶颈的第二刀",
             1.05, 5.62, 11.1, 0.44, 15, WHITE, True, PP_ALIGN.CENTER)

    s = base_slide(prs, "收端 Incast 绕路：换入口 NIC，保持 QP 内有序", "PXN data path")
    add_text(s, "正常路径", 0.8, 1.45, 1.2, 0.32, 14, CYAN, True)
    steps = [("GPU i", BLUE), ("NIC i", BLUE), ("Rail i", BLUE), ("NIC i", BLUE), ("GPU i", BLUE)]
    for i, (t, c) in enumerate(steps):
        x = 0.85 + i * 2.36
        pill(s, t, x, 1.95, 1.22, c, 13)
        if i < len(steps) - 1:
            line(s, x + 1.25, 2.14, x + 2.25, 2.14, MID, 3)
    add_text(s, "Incast 后下一 Epoch", 0.8, 3.18, 2.3, 0.32, 14, GREEN, True)
    steps = [("GPU i", BLUE), ("GPU k", GREEN), ("NIC k", GREEN), ("Rail k", GREEN), ("NIC k", GREEN), ("GPU k", GREEN), ("GPU i", BLUE)]
    for i, (t, c) in enumerate(steps):
        x = 0.82 + i * 1.78
        pill(s, t, x, 3.68, 0.98, c, 11)
        if i < len(steps) - 1:
            line(s, x + 1.0, 3.87, x + 1.68, 3.87, MID, 2)
    add_text(s, "NVLink/PXN", 1.73, 4.18, 1.0, 0.25, 9, CYAN, True, PP_ALIGN.CENTER)
    add_text(s, "Ethernet / ECMP（保序）", 5.68, 4.18, 2.15, 0.25, 9, CYAN, True, PP_ALIGN.CENTER)
    add_text(s, "NVLink/PXN", 10.49, 4.18, 1.0, 0.25, 9, CYAN, True, PP_ALIGN.CENTER)
    card(s, "容量判断", "CX8 800 Gb/s 约 100 GB/s；即使多个 Rail 向单 GPU 汇聚，NVLink5 仍有足够带宽。主要瓶颈仍是 Ethernet leaf，而不是 NVLink。", 0.85, 5.18, 11.55, 1.1, ORANGE, 14, 12)

    s = base_slide(prs, "遥测与判因：库内 Stall 为主，交换机 INT 为辅", "Observability")
    card(s, "通信库主信号", "每 VT：tx/rx FIFO stall\n完成字节、完成时间、重传/超时\n在线估计有效带宽与预计完成时间\n遥测 piggyback 在 collective 同步流量", 0.75, 1.52, 3.65, 3.3, BLUE)
    card(s, "交换机辅助信号", "QP 实际 ECMP 上行 path_id\nleaf→spine 出口队列 / ECN 档位\n目的 leaf 下行 / NIC 端口队列\n只观测与校验，不逐包改路", 4.84, 1.52, 3.65, 3.3, CYAN)
    card(s, "判因规则", "同 Rail 快慢分化 → L1 路径冲突\n同 Rail 全慢 + 下行热 → L2 Incast\n单 VT 连续超时 → 屏蔽\n探测恢复后逐步加入", 8.93, 1.52, 3.65, 3.3, GREEN)
    add_text(s, "没有交换机遥测时：用“同 Rail 是否全慢”作为启发式；准确性和收敛速度会下降",
             1.02, 5.55, 11.0, 0.48, 15, ORANGE, True, PP_ALIGN.CENTER)

    s = base_slide(prs, "Epoch 协议：只在 Collective 边界一致生效", "Consistency")
    phases = [
        ("1", "执行", "用 wₑ 完成本次 CCT", BLUE),
        ("2", "汇总", "交换同一 epoch 的 VT 遥测", CYAN),
        ("3", "计算", "各 rank 确定性生成 wₑ₊₁", GREEN),
        ("4", "生效", "下一次 CCT 同步切换", ORANGE),
    ]
    for i, (n, t, b, c) in enumerate(phases):
        x = 0.75 + i * 3.13
        rect(s, x, 1.65, 2.62, 2.25, DARK, True, c)
        pill(s, n, x + 0.18, 1.84, 0.46, c, 14)
        add_text(s, t, x + 0.78, 1.82, 1.35, 0.42, 18, WHITE, True)
        add_text(s, b, x + 0.2, 2.48, 2.2, 0.75, 13, LIGHT)
        if i < 3:
            line(s, x + 2.65, 2.78, x + 3.02, 2.78, MID, 3)
    card(s, "稳定性约束", "每 VT 保留最小探测份额  |  单步权重变化限幅  |  进入/退出阈值分离  |  部分故障屏蔽、周期探活  |  禁止迁移已提交到 QP 的字节",
         0.82, 4.62, 11.65, 1.0, RED, 14, 12)
    add_text(s, "核心不变量：所有 rank 对同一次 collective 使用完全相同的 VT 字节划分",
             1.05, 6.08, 11.0, 0.42, 16, CYAN, True, PP_ALIGN.CENTER)

    s = base_slide(prs, "分配器：层次化优化，不做独立 AIMD", "Allocator")
    rect(s, 0.8, 1.48, 12.0, 1.05, RGBColor(19, 49, 83), True)
    add_text(s, "目标：让所有健康 VT 的预计完成时间接近，同时最小化跨 Rail/PXN 成本",
             1.05, 1.74, 11.5, 0.45, 19, WHITE, True, PP_ALIGN.CENTER)
    for i, (t, b, c) in enumerate([
        ("Rail 内分配", "按 VT 有效带宽切份额\n保持 Rail 总份额不变", CYAN),
        ("Rail 间分配", "仅当 Rail 整体异常\n降低热 Rail 总份额", GREEN),
        ("约束投影", "最小份额、变化限幅\nNIC/GPU/PXN 容量上限", ORANGE),
    ]):
        card(s, t, b, 0.82 + i * 4.12, 3.05, 3.72, 1.9, c)
        if i < 2:
            line(s, 4.58 + i * 4.12, 4.0, 4.87 + i * 4.12, 4.0, MID, 3)
    add_text(s, "禁止：每 VT 独立 AIMD → 多个 VT 同时追逐“冷路径”，导致周期性振荡",
             1.1, 5.72, 11.0, 0.45, 16, RED, True, PP_ALIGN.CENTER)

    s = base_slide(prs, "实施边界与依赖", "Scope")
    card(s, "首期范围", "跨 NVL72 机柜的 Ethernet collective\n大消息 AllReduce / AllGather / RS\n再扩展 AlltoAll 与 EP", 0.75, 1.5, 3.7, 2.25, GREEN)
    card(s, "软件依赖", "通信库：独立 channel/QP 与 epoch\nRank 映射：4 GPU tray 对齐\nPXN：同 OS 内代理 GPU/NIC", 4.82, 1.5, 3.7, 2.25, BLUE)
    card(s, "网络依赖", "固定 sport 的 ECMP\n可查询/探测 path_id\nper-port 队列或 ECN/INT 遥测", 8.89, 1.5, 3.7, 2.25, CYAN)
    add_text(s, "明确不做", 0.78, 4.32, 1.4, 0.35, 15, RED, True)
    bullets(s, [
        "不优化机柜内 NVLink/NVLS；不把 72 张 NIC 做成单个 RDMA 端点池",
        "不逐包喷洒；不实现软件乱序接收或选择性重传",
        "不在 CCT 中途切权重；小消息默认关闭动态调整",
    ], 0.85, 4.72, 11.55, 1.75, 16)

    s = base_slide(prs, "验证路线与成功判据", "Validation")
    for i, (t, b, c) in enumerate([
        ("阶段 1\n路径与成本", "sport→path 稳定性\n1/2/4/8 QP 性能\n本地 NIC vs PXN", BLUE),
        ("阶段 2\n故障注入", "单上行冲突 → L1\n目的端 incast → L2\n限速/丢包/失效恢复", CYAN),
        ("阶段 3\n训练验证", "32/72/144/576 GPU\n多租户与多消息尺寸\nP50/P95/P99 CCT/JCT", GREEN),
    ]):
        card(s, t, b, 0.75 + i * 4.1, 1.5, 3.72, 2.58, c, 17, 12)
    rect(s, 0.8, 4.67, 12.0, 1.28, RGBColor(19, 49, 83), True)
    metrics = [
        ("≥98%", "无拥塞性能"),
        ("3–5 次", "L1 收敛"),
        ("显著下降", "Incast P95"),
        ("不挂死", "单 VT 故障"),
    ]
    for i, (v, l) in enumerate(metrics):
        x = 1.05 + i * 3.0
        add_text(s, v, x, 4.84, 1.6, 0.38, 20, WHITE, True, PP_ALIGN.CENTER)
        add_text(s, l, x, 5.28, 1.6, 0.28, 10, MID, False, PP_ALIGN.CENTER)

    s = base_slide(prs, "主要风险与应对", "Risk register")
    risks = [
        ("虚假多路径", "多个 QP 实际落同一上行", "初始化 path_id 校验；无正交则降 VT 数"),
        ("PXN 不可用", "Rank/并行维与 tray、rail 不对齐", "调度器按 4 GPU tray 对齐 communicator"),
        ("控制振荡", "遥测噪声、多个 VT 同时迁移", "层次化分配、滞回、限幅、全局目标"),
        ("资源反噬", "Channel 多导致小 chunk、QP cache/SM 开销", "大消息阈值；从 8 VT 起步，按收益扩展"),
        ("跨租户干扰", "借用兄弟 NIC 把拥塞转移给其它作业", "交换机队列辅助；限制跨 Rail/PXN 配额"),
    ]
    for i, (t, cause, mitigation) in enumerate(risks):
        y = 1.45 + i * 1.04
        pill(s, t, 0.78, y, 1.38, RED if i in (0, 2) else ORANGE, 11)
        add_text(s, cause, 2.42, y, 3.3, 0.4, 12, LIGHT)
        line(s, 5.8, y + 0.19, 6.35, y + 0.19, MID, 2)
        add_text(s, mitigation, 6.55, y - 0.02, 5.7, 0.46, 12, WHITE, True)

    s = base_slide(prs, "决策建议与下一步", "Recommendation")
    add_text(s, "建议立项：以 PReCCL 式字节重分配为首期，保留 Theseus 式调度热切换接口",
             0.82, 1.48, 11.7, 0.72, 22, WHITE, True, PP_ALIGN.CENTER)
    for i, (n, t, b, c) in enumerate([
        ("01", "固定架构", "8 VT / 4 Rail × 2 QP\nVT = Channel = QP", BLUE),
        ("02", "实现两级控制", "Rail 内 Path 优先\nIncast 才跨 Rail + PXN", CYAN),
        ("03", "交换机做观测", "path_id + 队列/ECN\n禁止逐包改路", GREEN),
        ("04", "用 P95 验收", "先做多租户大消息\n再扩展 AlltoAll/EP", ORANGE),
    ]):
        x = 0.72 + i * 3.13
        rect(s, x, 2.72, 2.75, 2.55, DARK, True, c)
        pill(s, n, x + 0.18, 2.92, 0.55, c, 12)
        add_text(s, t, x + 0.2, 3.58, 2.3, 0.38, 16, WHITE, True)
        add_text(s, b, x + 0.2, 4.12, 2.3, 0.72, 12, LIGHT)
    rect(s, 0.82, 5.9, 11.65, 0.6, RGBColor(19, 49, 83), True)
    add_text(s, "成功形态：稳态几乎无损，路径冲突先局部修复，Rail 级 Incast 再使用 NVLink 兜底",
             1.05, 5.99, 11.15, 0.38, 15, CYAN, True, PP_ALIGN.CENTER)

    prs.core_properties.title = "NVL72 多 QP 动态 VT 调度方案"
    prs.core_properties.subject = "PReCCL 式运行时字节重分配与跨 Rail PXN 绕路"
    prs.core_properties.author = "Cursor"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
