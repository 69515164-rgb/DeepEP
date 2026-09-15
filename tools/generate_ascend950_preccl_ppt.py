#!/usr/bin/env python3
"""Generate the Ascend 950 PReCCL design deck, including layered sequence slides."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap, qn
from pptx.util import Emu, Inches, Pt
from lxml import etree


OUT = Path(__file__).resolve().parents[1] / "docs" / "昇腾950_PReCCL运行时切分方案.pptx"
FIG = Path(__file__).resolve().parents[1] / "docs" / "figures"

NAVY = RGBColor(10, 24, 48)
BLUE = RGBColor(38, 108, 255)
CYAN = RGBColor(37, 200, 221)
GREEN = RGBColor(33, 181, 115)
ORANGE = RGBColor(255, 153, 51)
RED = RGBColor(225, 72, 72)
WHITE = RGBColor(248, 250, 253)
LIGHT = RGBColor(229, 236, 245)
MID = RGBColor(130, 148, 174)
DARK = RGBColor(28, 39, 57)
PALE = RGBColor(19, 49, 83)
HOP_I = RGBColor(32, 58, 92)
HOP_M = RGBColor(72, 54, 22)
HOP_J = RGBColor(24, 62, 42)
FONT = "Droid Sans Fallback"


def add_text(slide, text, x, y, w, h, size=20, color=WHITE, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line or fill
    return s


def connector(slide, x1, y1, x2, y2, color=MID, width=1.5, dashed=False):
    s = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2),
    )
    s.line.color.rgb = color
    s.line.width = Pt(width)
    if dashed:
        ln = s.line._get_or_add_ln()
        prst = etree.SubElement(ln, qn("a:prstDash"))
        prst.set("val", "dash")
    return s


def pill(slide, text, x, y, w, color=BLUE, size=12):
    rect(slide, x, y, w, 0.36, color, True)
    add_text(slide, text, x, y, w, 0.36, size, WHITE, True, PP_ALIGN.CENTER)


def card(slide, title, body, x, y, w, h, accent=BLUE, title_size=15, body_size=12):
    rect(slide, x, y, w, h, DARK, True, RGBColor(51, 69, 96))
    rect(slide, x, y, 0.07, h, accent)
    add_text(slide, title, x + 0.2, y + 0.1, w - 0.32, 0.32, title_size, WHITE, True)
    add_text(slide, body, x + 0.2, y + 0.46, w - 0.36, h - 0.58,
             body_size, LIGHT, False, valign=MSO_ANCHOR.TOP)


def bullets(slide, items, x, y, w, h, size=15, color=LIGHT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
    return box


def base_slide(prs, title, kicker=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    if kicker:
        add_text(slide, kicker, 0.6, 0.26, 8, 0.24, 10, CYAN, True)
    add_text(slide, title, 0.6, 0.5, 12.1, 0.5, 24, WHITE, True)
    rect(slide, 0.6, 1.06, 1.0, 0.045, BLUE)
    add_text(slide, str(len(prs.slides)), 12.3, 7.1, 0.45, 0.2, 8, MID, align=PP_ALIGN.RIGHT)
    return slide


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 13.333, 0.12, BLUE)
    pill(slide, "昇腾 950  /  HCCL", 0.7, 0.58, 2.15, CYAN, 12)
    add_text(slide, "PReCCL 式运行时拥塞检测\n与数据切分方案", 0.7, 1.2, 8.2, 1.7, 32, WHITE, True)
    add_text(slide, "AI CPU 模式为主：不换算法，只在算法拍边界改 tile→VT。检测在 L2 poll，切分在拍间 L2 Allocator。",
             0.72, 3.15, 7.6, 0.85, 16, LIGHT)
    rect(slide, 8.7, 1.15, 3.85, 4.55, DARK, True, RGBColor(51, 69, 96))
    for i, (name, col) in enumerate([
        ("L2 检测 stall", BLUE),
        ("拍间改切分", ORANGE),
        ("下一拍生效", GREEN),
    ]):
        pill(slide, name, 9.15, 1.7 + i * 1.15, 2.95, col, 14)
    add_text(slide, "时序图源文件\ndocs/figures/*.mmd  *.svg", 9.15, 5.05, 3.0, 0.5, 12, MID, False, PP_ALIGN.CENTER)
    add_text(slide, "技术方案  |  2026-09  |  对应 docs/ascend950_preccl_design.md", 0.72, 6.75, 8.5, 0.28, 12, MID)


def seq_lifelines(slide, actors, y0=1.28, y1=6.85):
    xs = []
    n = len(actors)
    left, right = 0.55, 12.75
    span = right - left
    for i, (name, col) in enumerate(actors):
        x = left + span * i / max(n - 1, 1)
        xs.append(x)
        connector(slide, x, y0 + 0.32, x, y1, RGBColor(70, 90, 120), 1.0)
        pill(slide, name, x - 0.62, y0, 1.24, col, 9)
    return xs


def seq_msg(slide, xs, i, j, y, text, color=LIGHT, dashed=False, size=10):
    x1, x2 = xs[i], xs[j]
    if abs(i - j) < 1e-6:
        connector(slide, x1, y, x1 + 0.55, y, CYAN, 1.2)
        add_text(slide, text, x1 + 0.6, y - 0.16, 2.4, 0.3, size, color)
        return
    connector(slide, min(x1, x2), y, max(x1, x2), y, CYAN if not dashed else MID, 1.3, dashed)
    mx = (x1 + x2) / 2
    add_text(slide, text, mx - 1.35, y - 0.22, 2.7, 0.22, size, color, False, PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_slide(prs)

    s = base_slide(prs, "结论：同一次 AllReduce 内，拍 i 检测，拍 i+1 生效", "Executive")
    card(s, "改什么", "只改下一拍各 VT 的连续 Tile 份额\n慢 VT 尾部划给快 VT\n总字节守恒", 0.6, 1.4, 3.9, 2.15, GREEN)
    card(s, "不改什么", "HCCL_ALGO / Ring 邻居 / HD 对端序列\n拍数、VT↔Jetty、已发 WQE\n算法边与角色", 4.7, 1.4, 3.9, 2.15, RED)
    card(s, "在哪一层改", "检测：L2 AI CPU poll stall\n改表：拍间 L2 Allocator\n生效：下一拍 L2 再下 WQE\nL0/L1 本集体不写表", 8.8, 1.4, 3.9, 2.15, ORANGE)
    rect(s, 0.6, 3.85, 12.15, 2.55, PALE, True)
    add_text(s, "AI CPU 拍循环插入点", 0.85, 4.0, 4, 0.32, 14, CYAN, True)
    add_text(s, "for step in 0 .. S-1:\n    按当前 tile→VT 给本拍各 VT 下 WQE     # L2→L6 传数\n    等本拍全部 CQE，记 stall                 # 同拍内不改表\n    用本拍 stall 算下一拍切分               # 唯一调整点",
             0.9, 4.38, 11.5, 1.8, 16, LIGHT, valign=MSO_ANCHOR.TOP)

    s = base_slide(prs, "两种展开模式：检测与切分对象不同", "Modes")
    card(s, "AI CPU / AICPU_TS  默认", "编排者：片上 AI CPU → STARS\nVT = 并发上下文 = 1 Jetty\nUB 4～8 VT，UBoE 4 VT\nstall 记在 AI CPU poll\n主战场：UBoE、大消息、回退、拍内切分",
         0.6, 1.4, 6.0, 3.35, BLUE, 16, 14)
    card(s, "CCU_SCHED  显式打开", "编排者：CCU → UB WQE\nVT = CCU Mission = 1 Jetty\nUB 8～16 VT，不用 CcuBuffer\n大 Reduce 会回退 AI CPU\n拍边界不一定暴露，逐拍切分后做",
         6.85, 1.4, 5.9, 3.35, CYAN, 16, 14)
    add_text(s, "禁止把两种 stall 计数直接比。T_hat 公式相同，必须用本模式自己的 poll 标定。",
             0.7, 5.05, 12, 0.4, 15, ORANGE, True)
    bullets(s, [
        "硬约束：VT = 一条可独立等待的 FIFO（Jetty / Transport Channel），禁止一条上下文喷多 Jetty",
        "CCU_MS / AIV / DeepEP P2P 不做动态切分；950PR 的 L2 只允许 IO Die 转发，禁止 HBM 中继",
    ], 0.7, 5.5, 12, 1.1, 14)

    s = base_slide(prs, "AI CPU 七层：上面编排，下面传数", "Layers")
    layers = [
        ("L0", "Host 控制面", "本集体不写表；跨集体才跑 Allocator", MID),
        ("L1", "HCCL API / AIC", "整次集体 Commit 一次、Wait 一次", BLUE),
        ("L2", "AI CPU 编排面", "按拍循环 / poll stall / 拍间改 tile→VT", ORANGE),
        ("L3", "STARS 调度面", "每拍每 VT 一条 TS 任务", CYAN),
        ("L4", "URMA 传输面", "Jetty k FIFO，已发 WQE 不改绑", GREEN),
        ("L5", "网络面", "Channel k → UB Port 或 UBoE", BLUE),
        ("L6", "本拍对端", "同 VT 收齐；尾包对齐下一拍长度表", GREEN),
    ]
    for i, (lv, name, desc, col) in enumerate(layers):
        y = 1.28 + i * 0.78
        pill(s, lv, 0.65, y, 0.7, col, 12)
        add_text(s, name, 1.5, y, 2.6, 0.36, 16, WHITE, True)
        add_text(s, desc, 4.2, y, 8.4, 0.36, 15, LIGHT)
    add_text(s, "热路径 L2→L6；切分只发生在黄块的 L2。L0 不在拍间。", 0.7, 6.85, 11, 0.28, 13, CYAN, True)

    s = base_slide(prs, "对象模型：一条 VT 一条 Jetty", "Binding")
    card(s, "硬绑定", "VT = AI CPU 并发上下文\n    = URMA Jetty\n    = 1 条 Transport Channel", 0.6, 1.4, 4.0, 2.4, BLUE, 16, 15)
    card(s, "UB 域", "4～8 VT\n4 个 Port 组 × 至多 2 Channel\n同组必须是独立 FIFO", 4.8, 1.4, 3.85, 2.4, CYAN, 16, 15)
    card(s, "UBoE 域", "4 VT = 2×400G × 2 Jetty\n不同 UDP 源端口做 ECMP 熵\n最像 NCCL channel=QP", 8.85, 1.4, 3.85, 2.4, GREEN, 16, 15)
    card(s, "切分单位", "用户 buffer → rank 块 → Tile → 恰好一个 VT\nTile：UB 1～4MB，UBoE 128～512KB\n迁徙只迁尾部连续区间", 0.6, 4.1, 6.0, 2.15, ORANGE, 15, 13)
    card(s, "检测主信号", "sq_stall / cq_stall / bytes_done 进 T_hat\nstream_stall 只辅信号，单独涨则冻权重\nT_hat = alpha + bytes_remain / B_hat", 6.85, 4.1, 5.85, 2.15, GREEN, 15, 13)

    s = base_slide(prs, "时序图文件：源文件与幻灯片对应", "Artifacts")
    rows = [
        ("图 A", "a_component_layers.mmd / .svg", "七层组件与控制/数据边"),
        ("图 D1", "d1_hop_layers.mmd / .svg", "三拍 × 七层（拍 i / 拍间 / 拍 i+1）"),
        ("图 D2", "d2_hop_sequence.mmd / .svg", "同一次 AllReduce 内完整拍间时序"),
        ("图 C", "c_epoch_sequence.mmd / .svg", "跨集体对照：Wait 后 L0 改表"),
    ]
    add_text(s, "目录  docs/figures/", 0.7, 1.28, 4, 0.3, 14, CYAN, True)
    for i, (g, f, d) in enumerate(rows):
        y = 1.7 + i * 0.85
        rect(s, 0.65, y, 12.05, 0.75, DARK, True, RGBColor(51, 69, 96))
        pill(s, g, 0.85, y + 0.18, 1.15, BLUE if i < 3 else MID, 12)
        add_text(s, f, 2.2, y + 0.08, 5.6, 0.28, 14, WHITE, True)
        add_text(s, d, 2.2, y + 0.38, 10, 0.26, 13, LIGHT)
    add_text(s, "后面三页把 D2 按时序拆成拍 i / 拍间 / 拍 i+1，层次与源文件一致。",
             0.7, 5.25, 12, 0.35, 15, ORANGE, True)
    bullets(s, [
        "设计正文：docs/ascend950_preccl_design.md 第 3.5 / 3.6 节",
        "重生成本 PPT：python3 tools/generate_ascend950_preccl_ppt.py",
    ], 0.7, 5.7, 12, 1.0, 14)

    # D2 hop i
    s = base_slide(prs, "图 D2-1  算法拍 i：L2→L6 只用 tile→VT_i", "Sequence  /  docs/figures/d2_hop_sequence.mmd")
    actors = [
        ("框架", MID), ("L0 Host", MID), ("L1 AIC", BLUE), ("L2 AI CPU", ORANGE),
        ("L3 STARS", CYAN), ("L4 Jetty", GREEN), ("L5 网络", BLUE), ("L6 对端", GREEN),
    ]
    rect(s, 0.5, 1.22, 12.35, 5.55, HOP_I, True)
    xs = seq_lifelines(s, actors, 1.35, 6.55)
    add_text(s, "算法图冻结：谁在哪一拍跟谁说话不改    L0 本集体不再写表", 0.7, 1.78, 12, 0.24, 11, CYAN, True, PP_ALIGN.CENTER)
    rows = [
        (2, 3, "Commit（整次集体只一次）", False),
        (3, 3, "按原算法挂本拍 chunk", False),
        (3, 4, "每 VT 一条 TS 任务", False),
        (4, 5, "doorbell SQE", False),
        (5, 6, "Transport Channel k", False),
        (6, 7, "本拍数据", False),
        (7, 5, "Notify", True),
        (5, 4, "CQE", True),
        (4, 3, "本拍该 VT 完成", True),
        (3, 3, "poll stall，已发 WQE 不改", False),
    ]
    for n, (a, b, t, dash) in enumerate(rows):
        seq_msg(s, xs, a, b, 2.15 + n * 0.42, t, LIGHT, dash, 10)
    add_text(s, "蓝块：只记账，不改表", 10.4, 6.6, 2.3, 0.25, 11, CYAN, True)

    # inter-hop
    s = base_slide(prs, "图 D2-2  拍间屏障：切分只在 L2", "Sequence  /  同一次 AllReduce 内")
    rect(s, 0.5, 1.22, 12.35, 5.55, HOP_M, True)
    xs = seq_lifelines(s, actors, 1.35, 6.55)
    add_text(s, "本拍全部 CQE 已收齐。L0 / L1 不参与。L3～L6 只捎带下一拍长度表。",
             0.7, 1.78, 12, 0.24, 12, ORANGE, True, PP_ALIGN.CENTER)
    rows = [
        (3, 4, "尾包捎带 w_i+1", False),
        (4, 5, "带内 footer（不新建连接）", False),
        (5, 6, "原 Channel 带出", False),
        (6, 7, "下一拍长度表", False),
        (7, 3, "对端对齐 / 同一公式", True),
        (3, 3, "Allocator 写 tile→VT_i+1", False),
    ]
    for n, (a, b, t, dash) in enumerate(rows):
        seq_msg(s, xs, a, b, 2.3 + n * 0.58, t, WHITE, dash, 12)
    rect(s, 3.6, 5.85, 6.2, 0.7, RGBColor(90, 60, 18), True)
    add_text(s, "唯一调整点：L2 写下一拍连续 Tile，不碰算法、不回 L0",
             3.7, 5.95, 6.0, 0.5, 13, WHITE, True, PP_ALIGN.CENTER)

    # hop i+1
    s = base_slide(prs, "图 D2-3  算法拍 i+1：同一 Jetty，新区间", "Sequence  /  生效点")
    rect(s, 0.5, 1.22, 12.35, 5.55, HOP_J, True)
    xs = seq_lifelines(s, actors, 1.35, 6.55)
    add_text(s, "仍走 L2→L6。Jetty / Channel / 对端角色不变，只是各口字节多少变。",
             0.7, 1.78, 12, 0.24, 12, GREEN, True, PP_ALIGN.CENTER)
    rows = [
        (3, 4, "按 tile→VT_i+1 下 Tile", False),
        (4, 5, "仍是 Jetty k", False),
        (5, 6, "仍是 Channel k", False),
        (6, 7, "角色不变，字节多少变", False),
        (7, 5, "Notify", True),
        (5, 3, "拍 i+1 CQE", True),
        (3, 2, "全部拍结束", False),
        (2, 0, "Wait 返回", True),
    ]
    for n, (a, b, t, dash) in enumerate(rows):
        seq_msg(s, xs, a, b, 2.2 + n * 0.5, t, LIGHT, dash, 11)
    add_text(s, "绿块：生效    L0 仍不改表", 10.2, 6.6, 2.5, 0.25, 11, GREEN, True)

    # D1 grid
    s = base_slide(prs, "图 D1  三拍 × 七层", "Layers × time  /  docs/figures/d1_hop_layers.mmd")
    cols = [
        (0.55, "入口", PALE, [("L0", "不写表"), ("L1", "Commit 一次")]),
        (2.85, "拍 i  蓝", HOP_I, [
            ("L2", "按 VT_i 挂 chunk"), ("L3", "每 VT 一条 TS"),
            ("L4", "Jetty doorbell"), ("L5", "Channel 出端口"),
            ("L6", "对端收齐"), ("L2", "poll stall"),
        ]),
        (5.55, "拍间  黄  唯一改表", HOP_M, [
            ("L0", "不参与"), ("L1", "不参与"),
            ("L2", "Allocator 写 VT_i+1"), ("L3-6", "只捎带 w_i+1"),
        ]),
        (8.25, "拍 i+1  绿", HOP_J, [
            ("L2", "按新表挂 Tile"), ("L3", "仍下原 VT"),
            ("L4", "仍是 Jetty k"), ("L5", "仍是 Channel k"),
            ("L6", "角色不变 字节变"),
        ]),
        (10.95, "出口", PALE, [("L1", "Wait 一次"), ("L0", "仍不改表")]),
    ]
    for x, title, col, cells in cols:
        w = 2.15 if x < 10 else 1.85
        rect(s, x, 1.28, w, 5.5, col, True)
        add_text(s, title, x + 0.06, 1.35, w - 0.1, 0.45, 12, WHITE, True, PP_ALIGN.CENTER)
        for i, (lv, txt) in enumerate(cells):
            y = 1.9 + i * 0.72
            pill(s, lv, x + 0.12, y, 0.55, BLUE if "L2" in lv and "拍间" in title else CYAN, 9)
            add_text(s, txt, x + 0.72, y, w - 0.85, 0.36, 11, LIGHT)

    s = base_slide(prs, "图 C 对照：跨集体才走 L0 Allocator", "Optional path  /  docs/figures/c_epoch_sequence.mmd")
    card(s, "拍内路径（推荐，AI CPU）", "一次 HcclAllReduce 内\n拍 i 收齐 → L2 改表 → 拍 i+1 生效\nL0 不写表", 0.6, 1.4, 6.0, 2.5, ORANGE, 16, 15)
    card(s, "跨集体路径（小消息 / CCU）", "CCT e Wait 之后\nL0 小 AllReduce + Allocator\n下一次 Prepare 读新表", 6.85, 1.4, 5.9, 2.5, BLUE, 16, 15)
    bullets(s, [
        "不要把「训练 step」和「算法拍」混为一谈。本方案主路径是算法拍。",
        "同一步有多次集体时，跨集体路径要指定 epoch 集体，避免同一步切分前后不一致。",
        "图模式要把 tile→VT 做成间接表，否则重放吃旧展开；拍内路径不依赖 Host 再下发。",
    ], 0.75, 4.2, 12, 2.0, 16)

    s = base_slide(prs, "Ring 可逐拍切，HD 默认只在相变切", "Ring vs HD")
    card(s, "Ring", "相邻两拍对端不变（next/prev）\nstep i 的 stall 可直接切 step i+1\n大消息（单拍 ≥1～2MB）可逐拍\n收发双方尾包对齐长度表",
         0.6, 1.4, 6.0, 3.15, GREEN, 16, 15)
    card(s, "HD Halving-Doubling", "每拍换对端（rank XOR 2^i）\n对端 A 的 stall 不能当对端 B 的 VT 权重\n默认只在 RS→AG 相变切一次\nPort 级 L2 权重可以跨拍留下",
         6.85, 1.4, 5.9, 3.15, ORANGE, 16, 15)
    bullets(s, [
        "拍间必须收齐：有 hop pipeline 就关掉，或改成每 K 拍 / 只在相变切",
        "不要为对齐再做一次全局小 AllReduce，会比一拍还贵",
        "CCU_SCHED 拍边界不一定暴露，逐拍切分首期只做 AI CPU",
    ], 0.75, 4.85, 12, 1.7, 15)

    s = base_slide(prs, "落地顺序与验收", "Plan")
    for i, (t, b, c) in enumerate([
        ("A  可观测", "Mission/上下文绑 1 Jetty\n打 stall，验证 stream_stall\n单独涨时不切 Port", BLUE),
        ("B  相变一次", "RS→AG 之间切一次\nRing 与 HD 都能做\n仍在同一次 AllReduce 内", CYAN),
        ("C  大消息逐拍", "仅 Ring、单拍够长\n尾包对齐长度表\n关 hop 重叠", GREEN),
        ("D  生产约束", "图模式间接表\n与 OP_RETRY 共存\nVT 数给融合留余量", ORANGE),
    ]):
        card(s, t, b, 0.55 + i * 3.18, 1.4, 3.0, 2.7, c, 15, 13)
    rect(s, 0.55, 4.4, 12.2, 2.15, PALE, True)
    add_text(s, "验收", 0.8, 4.55, 1.2, 0.3, 14, CYAN, True)
    bullets(s, [
        "限速 1 个 UB Port 或 1 条 UBoE 到 10%：健康 VT 仍满速，集体不跟死链路走",
        "STARS / TS 轨迹：没有中途改 Jetty；权重只在拍边界或集体边界变",
        "各 rank 下一拍长度表一致；950PR 上 L2 无额外 HBM 中继拷贝",
    ], 0.8, 4.95, 11.7, 1.45, 14)

    s = base_slide(prs, "风险与明确不做", "Risks")
    risks = [
        ("拍重叠", "hop i+1 发送与 hop i 接收并行则没有屏障", "关 pipeline，或退回相变 / 跨集体"),
        ("HD 错切", "把旧对端 stall 用到新对端 VT", "HD 不默认逐拍，只切 RS→AG"),
        ("cache / 图", "AI CPU cache 或 aclgraph 重放旧表", "CacheDisable；tile→VT 间接读"),
        ("stall 混用", "AI CPU 与 CCU 计数直接比较", "T_hat 必须用本模式标定"),
        ("CTP 丢包", "无端到端 SACK，不能 Jetty 内重排", "慢 VT 降权或隔离，走 OP_RETRY"),
    ]
    for i, (t, cause, fix) in enumerate(risks):
        y = 1.32 + i * 1.05
        pill(s, t, 0.65, y, 1.45, RED if i in (0, 1) else ORANGE, 11)
        add_text(s, cause, 2.3, y, 5.1, 0.4, 13, LIGHT)
        add_text(s, fix, 7.5, y, 5.2, 0.4, 13, WHITE, True)

    prs.core_properties.title = "昇腾950 PReCCL式运行时拥塞检测与数据切分"
    prs.core_properties.subject = "AI CPU 模式算法拍时序与层次"
    prs.core_properties.author = "Cursor"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
